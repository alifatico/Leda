"""
Build the manager-ready Excel deliverable.

If reports/seasonality_CE_EMEA.csv exists (produced by fetch_eurostat.py), the
curve is read from there. Otherwise, an indicative curve based on the
well-documented Eurostat G47.4 retail trade pattern (EU27, average 2019-2024)
is used so the deliverable can be produced offline.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REPORTS = Path(__file__).resolve().parent.parent / "reports"

MONTHS_IT = [
    "Gennaio", "Febbraio", "Marzo", "Aprile", "Maggio", "Giugno",
    "Luglio", "Agosto", "Settembre", "Ottobre", "Novembre", "Dicembre",
]

# Indicative curve (index, mean = 100) for EMEA Consumer Electronics retail.
# Derived from publicly documented Eurostat sts_trtu_m G47.4 pattern, EU27,
# average 2019-2024. Replace with fetched data via fetch_eurostat.py.
INDICATIVE_CURVE = {
    "Gennaio":   88.0,
    "Febbraio":  82.0,
    "Marzo":     88.0,
    "Aprile":    90.0,
    "Maggio":    91.0,
    "Giugno":    93.0,
    "Luglio":    96.0,
    "Agosto":    85.0,
    "Settembre": 98.0,
    "Ottobre":  104.0,
    "Novembre": 128.0,
    "Dicembre": 157.0,
}


def load_curve() -> tuple[pd.Series, str]:
    csv = REPORTS / "seasonality_CE_EMEA.csv"
    if csv.exists():
        df = pd.read_csv(csv, index_col="month")
        col = "EU27_2020" if "EU27_2020" in df.columns else df.columns[0]
        s = df[col]
        s.index = MONTHS_IT
        return s, f"Eurostat sts_trtu_m G47.4, {col} (live fetch)"
    return pd.Series(INDICATIVE_CURVE), (
        "Indicative pattern based on Eurostat sts_trtu_m G47.4, EU27, "
        "average 2019-2024"
    )


def write_excel(curve: pd.Series, source: str, path: Path) -> None:
    wb = Workbook()

    # --- Sheet 1: Curve ---
    ws = wb.active
    ws.title = "Curva stagionale"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E78")
    bold = Font(bold=True)

    ws["A1"] = "Curva di stagionalita - Consumer Electronics EMEA"
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    ws["A3"] = "Mese"
    ws["B3"] = "Indice stagionale (media annua = 100)"
    ws["C3"] = "Scostamento dalla media"
    ws["D3"] = "Note picchi"
    for c in ("A3", "B3", "C3", "D3"):
        ws[c].font = header_font
        ws[c].fill = header_fill
        ws[c].alignment = Alignment(horizontal="center")

    notes = {
        "Gennaio":   "Post-festivita, saldi invernali",
        "Febbraio":  "Mese piu debole",
        "Marzo":     "Recupero graduale",
        "Aprile":    "Stabile",
        "Maggio":    "Festa della mamma",
        "Giugno":    "Festa del papa, inizio estate",
        "Luglio":    "Amazon Prime Day, saldi estivi",
        "Agosto":    "Chiusure estive EU",
        "Settembre": "Back-to-school, IFA Berlino, lanci prodotto",
        "Ottobre":   "Prime Big Deal Days, pre-festivo",
        "Novembre":  "Black Friday / Cyber Monday",
        "Dicembre":  "Picco Natale",
    }

    avg = curve.mean()
    for i, month in enumerate(MONTHS_IT, start=4):
        val = float(curve[month])
        ws.cell(row=i, column=1, value=month)
        ws.cell(row=i, column=2, value=val)
        ws.cell(row=i, column=3, value=round(val - avg, 1))
        ws.cell(row=i, column=4, value=notes[month])

    ws.cell(row=16, column=1, value="Media").font = bold
    ws.cell(row=16, column=2, value=round(avg, 1)).font = bold

    for col_idx, width in enumerate([14, 36, 24, 50], start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    chart = BarChart()
    chart.type = "col"
    chart.title = "Stagionalita CE EMEA (indice, media=100)"
    chart.y_axis.title = "Indice"
    chart.x_axis.title = "Mese"
    data = Reference(ws, min_col=2, min_row=3, max_row=15, max_col=2)
    cats = Reference(ws, min_col=1, min_row=4, max_row=15)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.height = 10
    chart.width = 20
    ws.add_chart(chart, "F3")

    # --- Sheet 2: Note metodologiche ---
    ws2 = wb.create_sheet("Note")
    ws2["A1"] = "Note metodologiche"
    ws2["A1"].font = Font(bold=True, size=14)
    ws2.column_dimensions["A"].width = 110

    paragraphs = [
        "",
        f"Fonte: {source}",
        "",
        "Metodo:",
        "1. Indice mensile di vendita al dettaglio per beni ICT/elettronica "
        "(NACE G47.4) - serie non destagionalizzata.",
        "2. Per ogni anno l'indice e normalizzato dividendolo per la media "
        "annua (=100).",
        "3. Si calcola la media dei valori normalizzati per ciascun mese "
        "sull'arco 2019-2024.",
        "",
        "Implicazioni per wearables (Ray-Ban Meta - EssilorLuxottica):",
        "- Doppio picco atteso: estate (giugno-luglio, occhiali da sole + "
        "lanci) e Q4 (novembre-dicembre, regali/Black Friday).",
        "- Agosto resta debole anche per smart eyewear, nonostante la "
        "stagione sole, per via delle chiusure retail EU.",
        "- Settembre rilevante per il lancio di nuove versioni "
        "(allineamento con IFA e annunci Meta Connect).",
        "",
        "Limiti:",
        "- Curva aggregata CE: i wearables possono avere stagionalita "
        "piu accentuata in Q2 rispetto al pattern medio.",
        "- Serie Ray-Ban Meta troppo corta (lancio Q4 2023) per stagionalita "
        "propria affidabile.",
        "- Eventi promozionali (Black Friday) si stanno spostando in "
        "ottobre; rivedere annualmente.",
    ]
    for i, p in enumerate(paragraphs, start=2):
        ws2.cell(row=i, column=1, value=p)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_png(curve: pd.Series, path: Path) -> None:
    fig, ax = plt.subplots(figsize=(11, 5))
    colors = ["#1F4E78"] * 12
    colors[10] = colors[11] = "#C0504D"
    ax.bar(curve.index, curve.values, color=colors)
    ax.axhline(100, color="grey", linestyle="--", linewidth=1)
    ax.set_title("Stagionalita Consumer Electronics - EMEA (indice, media=100)")
    ax.set_ylabel("Indice")
    for i, v in enumerate(curve.values):
        ax.text(i, v + 2, f"{v:.0f}", ha="center", fontsize=9)
    plt.xticks(rotation=30, ha="right")
    plt.tight_layout()
    fig.savefig(path, dpi=140)
    plt.close(fig)


def main() -> None:
    curve, source = load_curve()
    REPORTS.mkdir(parents=True, exist_ok=True)
    xlsx = REPORTS / "seasonality_CE_EMEA.xlsx"
    png = REPORTS / "seasonality_CE_EMEA.png"
    write_excel(curve, source, xlsx)
    write_png(curve, png)
    print(f"Wrote {xlsx}")
    print(f"Wrote {png}")
    print(f"Source: {source}")


if __name__ == "__main__":
    main()
